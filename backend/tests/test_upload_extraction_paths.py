import io
import os
import sys
import unittest

from fastapi import HTTPException, UploadFile
from PIL import Image, ImageDraw, ImageFont

os.environ.setdefault("SECRET_KEY", "test-secret-key")
os.environ.setdefault("DATABASE_URL", "sqlite://")
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

import models  # noqa: E402,F401
from routers import ingest  # noqa: E402


def upload(filename: str, content: bytes, mime: str) -> UploadFile:
    return UploadFile(filename=filename, file=io.BytesIO(content), headers={"content-type": mime})


def readable_image() -> bytes:
    image = Image.new("RGB", (1800, 1000), "white")
    draw = ImageDraw.Draw(image)
    font_path = r"C:\Windows\Fonts\arial.ttf"
    font = ImageFont.truetype(font_path, 54) if os.path.exists(font_path) else None
    lines = [
        "CSC 301 SYSTEM ANALYSIS AND DESIGN",
        "Lecture note: requirements, architecture, and verification",
        "Study the system boundary, stakeholders, and test evidence.",
        "Review this material before the examination.",
    ]
    for index, line in enumerate(lines):
        draw.text((90, 100 + index * 170), line, fill="black", font=font)
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


class UploadExtractionPathTests(unittest.TestCase):
    def test_pdf_docx_and_pptx_text_are_extracted(self):
        import fitz
        from docx import Document
        from pptx import Presentation

        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 90), "CSC 301 lecture notes: architecture and verification evidence for an examination.")
        pdf_bytes = pdf.tobytes()
        pdf.close()

        document = Document()
        document.add_paragraph("CSC 301 lecture notes: requirements, system architecture, verification, stakeholders, and examination revision guidance.")
        docx_buffer = io.BytesIO()
        document.save(docx_buffer)

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "CSC 301 Revision"
        slide.placeholders[1].text = "Architecture, requirements, verification, and examination practice."
        pptx_buffer = io.BytesIO()
        presentation.save(pptx_buffer)

        cases = [
            ("notes.pdf", pdf_bytes, "application/pdf", "embedded_text"),
            ("notes.docx", docx_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx_text"),
            ("notes.pptx", pptx_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation", "pptx_text"),
        ]
        for filename, content, mime, expected_method in cases:
            with self.subTest(filename=filename):
                result = ingest.extract_pdf_text(upload(filename, content, mime))
                self.assertEqual(result["method"], expected_method)
                self.assertTrue(any(term in result["text"].lower() for term in ("requirements", "architecture", "verification")))
                self.assertTrue(result["cleaned_text"].strip())
                self.assertEqual(result["file_bytes"], content)

    def test_printed_image_and_scanned_pdf_use_real_ocr(self):
        import fitz

        image_bytes = readable_image()
        image_result = ingest.extract_pdf_text(upload("printed.png", image_bytes, "image/png"))
        self.assertEqual(image_result["method"], "ocr")
        self.assertTrue(image_result["ocr_used"])
        self.assertIn("architecture", image_result["text"].lower())
        self.assertTrue(image_result["searchable"])

        scanned_pdf = fitz.open()
        page = scanned_pdf.new_page(width=612, height=792)
        page.insert_image(fitz.Rect(30, 30, 582, 600), stream=image_bytes)
        scanned_bytes = scanned_pdf.tobytes()
        scanned_pdf.close()
        pdf_result = ingest.extract_pdf_text(upload("scanned.pdf", scanned_bytes, "application/pdf"))
        self.assertEqual(pdf_result["method"], "ocr")
        self.assertTrue(pdf_result["ocr_used"])
        self.assertIn("requirements", pdf_result["text"].lower())
        self.assertTrue(pdf_result["searchable"])

    def test_unsupported_format_is_rejected_before_extraction(self):
        with self.assertRaises(HTTPException) as context:
            ingest.extract_pdf_text(upload("notes.txt", b"plain text", "text/plain"))
        self.assertEqual(context.exception.status_code, 400)


if __name__ == "__main__":
    unittest.main()
